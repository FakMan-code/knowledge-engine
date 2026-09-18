"""Word, Excel and PowerPoint.

Each one is optional and independent: having ``openpyxl`` but not
``python-docx`` means spreadsheets are read and documents are reported as a
gap, rather than all Office files failing together.

Locators follow what the format itself addresses. A Word paragraph has no line
number, a spreadsheet row is ``Hoja1!A5:D5``, a slide is a slide. Pretending any
of them is a line range would produce citations that cannot be reopened.
"""

from __future__ import annotations

import io

from ingest.extractors.base import Extractor, MAX_CHARS
from ingest.models import EmptyExtraction, ExtractedUnit

from core.provenance import Cell, Node, Slide

try:
    import docx
except ImportError:  # optional dependency
    docx = None

try:
    import openpyxl
except ImportError:  # optional dependency
    openpyxl = None

try:
    import pptx
except ImportError:  # optional dependency
    pptx = None

ROWS_PER_UNIT = 50


class DocxExtractor(Extractor):
    name = "docx"
    media_types = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    extensions = (".docx",)
    requirement = "docx"
    install_hint = "pip install python-docx"

    def extract(self, data: bytes, path: str) -> list[ExtractedUnit]:
        self.require()
        document = docx.Document(io.BytesIO(data))

        units: list[ExtractedUnit] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = " ".join(paragraph.text.split())
            if not text:
                continue
            units.append(
                ExtractedUnit(
                    locator=Node(path=f"/document/paragraph[{index}]"),
                    text=text,
                    meta={"style": paragraph.style.name if paragraph.style else ""},
                    verbatim=False,
                )
            )

        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                if not any(cells):
                    continue
                units.append(
                    ExtractedUnit(
                        locator=Node(path=f"/document/table[{table_index}]/row[{row_index}]"),
                        text=" | ".join(cells),
                        meta={"table": table_index},
                        verbatim=False,
                    )
                )

        if not units:
            raise EmptyExtraction("docx_sin_texto", "el documento no tiene párrafos con texto")
        return units


class XlsxExtractor(Extractor):
    name = "xlsx"
    media_types = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",)
    extensions = (".xlsx", ".xlsm")
    requirement = "openpyxl"
    install_hint = "pip install openpyxl"

    def extract(self, data: bytes, path: str) -> list[ExtractedUnit]:
        self.require()
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)

        units: list[ExtractedUnit] = []
        for sheet in book.worksheets:
            batch: list[tuple[int, list[str]]] = []
            for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = ["" if value is None else str(value).strip() for value in row]
                if any(values):
                    batch.append((number, values))
                if len(batch) >= ROWS_PER_UNIT:
                    units.append(_sheet_unit(sheet.title, batch))
                    batch = []
            if batch:
                units.append(_sheet_unit(sheet.title, batch))
        book.close()

        if not units:
            raise EmptyExtraction("xlsx_sin_datos", "todas las hojas están vacías")
        return units


def _sheet_unit(sheet: str, batch: list[tuple[int, list[str]]]) -> ExtractedUnit:
    width = max(len(values) for _, values in batch)
    last_column = _column_letter(width)
    ref = f"A{batch[0][0]}:{last_column}{batch[-1][0]}"
    body = "\n".join(" | ".join(values) for _, values in batch)
    return ExtractedUnit(
        locator=Cell(sheet=sheet, ref=ref),
        text=body[:MAX_CHARS],
        meta={"rows": len(batch), "columns": width},
        verbatim=False,
    )


def _column_letter(index: int) -> str:
    letters = ""
    index = max(1, index)
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


class PptxExtractor(Extractor):
    name = "pptx"
    media_types = (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    extensions = (".pptx",)
    requirement = "pptx"
    install_hint = "pip install python-pptx"

    def extract(self, data: bytes, path: str) -> list[ExtractedUnit]:
        self.require()
        deck = pptx.Presentation(io.BytesIO(data))

        units: list[ExtractedUnit] = []
        for number, slide in enumerate(deck.slides, start=1):
            for shape in slide.shapes:
                text = " ".join(getattr(shape, "text", "").split())
                if not text:
                    continue
                units.append(
                    ExtractedUnit(
                        locator=Slide(number=number, shape=shape.name or None),
                        text=text,
                        meta={"slides": len(deck.slides._sldIdLst)},
                        verbatim=False,
                    )
                )

        if not units:
            raise EmptyExtraction("pptx_sin_texto", "ninguna diapositiva tiene texto")
        return units
